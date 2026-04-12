# app/utils/ranking_export.py
#
# Genera archivos descargables del ranking de un torneo.
# Funciones:
#   - ranking_a_excel(torneo, ranking)  → BytesIO con .xlsx
#   - ranking_a_pdf(torneo, ranking)    → BytesIO con .pdf

from __future__ import annotations

import io
from datetime import datetime

# ── Excel ────────────────────────────────────────────────────────────────────

def ranking_a_excel(torneo, ranking: list) -> io.BytesIO:
    """Devuelve un BytesIO con el Excel del ranking listo para enviar."""
    from openpyxl import Workbook
    from openpyxl.styles import (Alignment, Border, Font, PatternFill,
                                 Side)
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Ranking'

    # ── Paleta de colores ────────────────────────────────────────────────────
    AZUL_OSC  = '00003263'
    AZUL_MED  = '000D47A1'
    AMARILLO  = '00FFF176'
    PLATA     = '00E0E0E0'
    BRONCE    = '00D7CCC8'
    BLANCO    = '00FFFFFF'
    GRIS_CLR  = '00F5F5F5'
    BORDE_CLR = '00BDBDBD'

    thin = Side(style='thin', color=BORDE_CLR)
    borde = Border(left=thin, right=thin, top=thin, bottom=thin)

    def fill(hex_color):
        return PatternFill('solid', fgColor=hex_color)

    # ── Título ───────────────────────────────────────────────────────────────
    ws.merge_cells('A1:F1')
    c = ws['A1']
    c.value        = f'Ranking — {torneo.nombre}'
    c.font         = Font(bold=True, size=14, color=BLANCO)
    c.fill         = fill(AZUL_OSC)
    c.alignment    = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    # ── Subtítulo / fecha ────────────────────────────────────────────────────
    ws.merge_cells('A2:F2')
    c = ws['A2']
    c.value     = f'Generado el {datetime.now().strftime("%d/%m/%Y %H:%M")}'
    c.font      = Font(italic=True, size=9, color='00757575')
    c.fill      = fill(AZUL_MED)
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.font      = Font(italic=True, size=9, color=BLANCO)
    ws.row_dimensions[2].height = 16

    # ── Encabezados ──────────────────────────────────────────────────────────
    headers = ['Pos.', 'Jugador', 'Sector', 'Pronosticados', 'Exactos 🎯', 'Puntos']
    for col_idx, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=col_idx, value=h)
        c.font      = Font(bold=True, size=10, color=BLANCO)
        c.fill      = fill(AZUL_OSC)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border    = borde
    ws.row_dimensions[3].height = 20

    # ── Filas de datos ───────────────────────────────────────────────────────
    MEDALLA = {1: '🥇', 2: '🥈', 3: '🥉'}

    for row_idx, r in enumerate(ranking, 4):
        pos      = r.get('posicion', row_idx - 3)
        jugador  = r.get('usuario', '')
        sector   = r.get('sector') or ''
        pron     = r.get('pronosticados', 0) or 0
        exactos  = r.get('exactos', 0) or 0
        puntos   = r.get('puntos', 0) or 0

        fila = [
            MEDALLA.get(pos, str(pos)),
            jugador,
            sector,
            pron,
            exactos,
            puntos,
        ]

        if pos == 1:
            row_fill = fill(AMARILLO)
        elif pos == 2:
            row_fill = fill(PLATA)
        elif pos == 3:
            row_fill = fill(BRONCE)
        elif row_idx % 2 == 0:
            row_fill = fill(GRIS_CLR)
        else:
            row_fill = fill(BLANCO)

        for col_idx, valor in enumerate(fila, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=valor)
            c.fill   = row_fill
            c.border = borde
            c.alignment = Alignment(
                horizontal='center' if col_idx != 2 else 'left',
                vertical='center',
            )
            if col_idx == 6:  # Puntos — negrita
                c.font = Font(bold=True, size=11)
        ws.row_dimensions[row_idx].height = 18

    # ── Ancho de columnas ────────────────────────────────────────────────────
    anchos = [7, 30, 18, 16, 12, 10]
    for i, ancho in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    # ── Congelar fila de encabezado ──────────────────────────────────────────
    ws.freeze_panes = 'A4'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── PDF ──────────────────────────────────────────────────────────────────────

def ranking_a_pdf(torneo, ranking: list) -> io.BytesIO:
    """Devuelve un BytesIO con el PDF del ranking listo para enviar."""
    from fpdf import FPDF

    AZUL_R, AZUL_G, AZUL_B           = 0,   50,  99   # #003263
    AZUL_MED_R, AZUL_MED_G, AZUL_MED_B = 13, 71, 161  # #0D47A1
    ORO_R, ORO_G, ORO_B              = 255, 213,  79   # oro
    PLATA_R, PLATA_G, PLATA_B        = 200, 200, 200
    BRONCE_R, BRONCE_G, BRONCE_B     = 188, 143, 143
    GRIS_R, GRIS_G, GRIS_B          = 245, 245, 245

    class RankingPDF(FPDF):
        def header(self):
            # Franja azul de título
            self.set_fill_color(AZUL_R, AZUL_G, AZUL_B)
            self.rect(0, 0, 210, 22, 'F')
            self.set_font('Helvetica', 'B', 14)
            self.set_text_color(255, 255, 255)
            self.set_xy(10, 6)
            nombre_raw = torneo.nombre[:60] + ('...' if len(torneo.nombre) > 60 else '')
            nombre = nombre_raw.encode('latin-1', errors='replace').decode('latin-1')
            self.cell(0, 10, f'Ranking - {nombre}', new_x='LMARGIN', new_y='NEXT')
            self.set_font('Helvetica', '', 8)
            self.set_text_color(200, 220, 255)
            self.set_x(10)
            self.cell(0, 5,
                      f'Generado el {datetime.now().strftime("%d/%m/%Y a las %H:%M")}',
                      new_x='LMARGIN', new_y='NEXT')
            self.ln(4)
            self.set_text_color(50, 50, 50)

        def footer(self):
            self.set_y(-12)
            self.set_font('Helvetica', 'I', 7)
            self.set_text_color(150, 150, 150)
            self.cell(0, 10,
                      f'Pagina {self.page_no()} | Prode Deportivo',
                      align='C')

    pdf = RankingPDF(orientation='P', unit='mm', format='A4')
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # ── Encabezado de tabla ──────────────────────────────────────────────────
    COL_W = [14, 72, 30, 24, 20, 20]  # anchos de columna en mm
    HDR   = ['Pos.', 'Jugador', 'Sector', 'Pronos.', 'Exactos', 'Pts']

    pdf.set_fill_color(AZUL_R, AZUL_G, AZUL_B)
    pdf.set_text_color(255, 255, 255)
    pdf.set_font('Helvetica', 'B', 8)
    for w, h in zip(COL_W, HDR):
        pdf.cell(w, 7, h, border=1, align='C', fill=True)
    pdf.ln()

    # ── Filas ────────────────────────────────────────────────────────────────
    MEDALLA_TXT = {1: '1ro', 2: '2do', 3: '3ro'}
    pdf.set_font('Helvetica', '', 8)

    for idx, r in enumerate(ranking):
        pos     = r.get('posicion', idx + 1)
        jugador_raw = r.get('usuario', '')[:35]
        jugador = jugador_raw.encode('latin-1', errors='replace').decode('latin-1')
        sector_raw  = (r.get('sector') or '')[:20]
        sector  = sector_raw.encode('latin-1', errors='replace').decode('latin-1')
        pron    = str(r.get('pronosticados', 0) or 0)
        exactos = str(r.get('exactos', 0) or 0)
        puntos  = str(r.get('puntos', 0) or 0)

        if pos == 1:
            pdf.set_fill_color(ORO_R, ORO_G, ORO_B)
            pdf.set_text_color(80, 60, 0)
            fill_row = True
        elif pos == 2:
            pdf.set_fill_color(PLATA_R, PLATA_G, PLATA_B)
            pdf.set_text_color(50, 50, 50)
            fill_row = True
        elif pos == 3:
            pdf.set_fill_color(BRONCE_R, BRONCE_G, BRONCE_B)
            pdf.set_text_color(60, 30, 30)
            fill_row = True
        elif idx % 2 == 0:
            pdf.set_fill_color(GRIS_R, GRIS_G, GRIS_B)
            pdf.set_text_color(50, 50, 50)
            fill_row = True
        else:
            pdf.set_fill_color(255, 255, 255)
            pdf.set_text_color(50, 50, 50)
            fill_row = True

        pos_txt = MEDALLA_TXT.get(pos, str(pos))
        fila = [pos_txt, jugador, sector, pron, exactos, puntos]

        for i, (w, val) in enumerate(zip(COL_W, fila)):
            is_puntos = (i == 5)
            if is_puntos:
                pdf.set_font('Helvetica', 'B', 9)
            alin = 'L' if i == 1 else 'C'
            pdf.cell(w, 6, val, border=1, align=alin, fill=fill_row)
            if is_puntos:
                pdf.set_font('Helvetica', '', 8)
        pdf.ln()

    # ── Totales / pie de tabla ───────────────────────────────────────────────
    pdf.ln(2)
    pdf.set_font('Helvetica', 'I', 7)
    pdf.set_text_color(130, 130, 130)
    pdf.set_fill_color(255, 255, 255)
    pdf.cell(0, 5, f'Total de participantes: {len(ranking)}', align='L')

    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return buf
